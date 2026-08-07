#!/usr/bin/env python3
"""넥스트웨이브 경영데이터 xlsx -> Next.js 정적 JSON ETL.

엑셀 수식(SUMIFS/DATEDIF/ROUND 등)을 파이썬으로 재현해 계산 결과까지 JSON에 담는다.
"""
import json
import os
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.environ.get(
    "XLSX_PATH",
    os.path.join(ROOT, "..", "바이브코딩_실습데이터", "넥스트웨이브_경영데이터_2026-07.xlsx"),
)
OUT = os.path.join(ROOT, "src", "data")

BASE_DATE = date(2026, 7, 31)
VAT_RATE = 0.1


def xround(x, digits=0):
    """Excel ROUND (half away from zero)."""
    if x is None:
        return 0
    q = Decimal(10) ** -digits
    return float(Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP))


def d(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    return v


def iso(v):
    v = d(v)
    return v.isoformat() if v else None


def rows(ws, header_row=1, last_row=None):
    hdr = [c.value for c in ws[header_row]]
    out = []
    end = last_row or ws.max_row
    for r in ws.iter_rows(min_row=header_row + 1, max_row=end, values_only=True):
        if r[0] is None:
            continue
        out.append(dict(zip(hdr, r)))
    return out


def months_between(start, end):
    """DATEDIF(start, end, "m")"""
    m = (end.year - start.year) * 12 + (end.month - start.month)
    if end.day < start.day:
        m -= 1
    return max(m, 0)


wb = openpyxl.load_workbook(SRC, data_only=False)

# ---------------------------------------------------------------- 01 매출상세
sales = []
for r in rows(wb["01_매출상세"], last_row=240):
    qty = r["수량"] or 0
    price = r["단가"] or 0
    supply = qty * price
    vat = xround(supply * VAT_RATE)
    sales.append(
        {
            "date": iso(r["일자"]),
            "voucherNo": r["전표번호"],
            "customerCode": r["거래처코드"],
            "customer": r["거래처명"],
            "industry": r["산업군"],
            "region": r["지역"],
            "productCode": r["제품코드"],
            "product": r["제품명"],
            "category": r["카테고리"],
            "qty": qty,
            "unitPrice": price,
            "supplyAmount": supply,
            "vat": vat,
            "total": supply + vat,
            "rep": r["영업담당자"],
            "team": r["부서"],
            "paymentTerm": r["결제조건"],
            "channel": r["유입채널"],
        }
    )

# ------------------------------------------------------------- 03 회계 전표
vouchers = []
for r in rows(wb["03_회계_전표"], last_row=133):
    vouchers.append(
        {
            "date": iso(r["전표일자"]),
            "voucherNo": r["전표번호"],
            "accountCode": r["계정코드"],
            "account": r["계정과목"],
            "accountType": r["계정구분"],
            "dept": r["부서"],
            "vendor": r["거래처"],
            "memo": r["적요"],
            "amount": r["금액"] or 0,
            "evidence": r["증빙유형"],
            "approval": r["결재상태"],
        }
    )

# ------------------------------------------------------------ 02 손익계산서
pl_ws = wb["02_손익계산서"]
prev_by_label = {}
pl_meta = []  # (구분, 계정코드, 계정과목, 전월)
for r in pl_ws.iter_rows(min_row=5, max_row=41, values_only=True):
    group, code, label, _cur, prev = r[0], r[1], r[2], r[3], r[4]
    if label:
        pl_meta.append((group, code, label, prev))
        prev_by_label[label] = prev


def sum_sales_by_category(cat):
    return sum(s["supplyAmount"] for s in sales if s["category"] == cat)


def sum_voucher_by_code(code):
    return sum(v["amount"] for v in vouchers if v["accountCode"] == code)


revenue_lines = [
    {"label": "SaaS구독 매출", "current": sum_sales_by_category("SaaS구독")},
    {"label": "하드웨어 매출", "current": sum_sales_by_category("하드웨어")},
    {"label": "용역 매출", "current": sum_sales_by_category("용역")},
    {"label": "유지보수 매출", "current": sum_sales_by_category("유지보수")},
]
cogs_lines = [
    {"code": c, "label": l, "current": sum_voucher_by_code(c)}
    for c, l in [(451, "하드웨어매입원가"), (452, "외주용역비"), (453, "클라우드/인프라비")]
]
sga_codes = [
    (511, "급여"), (512, "상여금"), (513, "퇴직급여충당금"), (514, "복리후생비"),
    (515, "여비교통비"), (516, "접대비"), (517, "통신비"), (518, "소모품비"),
    (519, "지급수수료"), (520, "광고선전비"), (521, "임차료"), (522, "감가상각비"),
    (523, "보험료"), (524, "교육훈련비"), (525, "세금과공과"), (526, "차량유지비"),
]
sga_lines = [{"code": c, "label": l, "current": sum_voucher_by_code(c)} for c, l in sga_codes]

for line in revenue_lines + cogs_lines + sga_lines:
    line["previous"] = prev_by_label.get(line["label"], 0) or 0

revenue = sum(x["current"] for x in revenue_lines)
revenue_prev = sum(x["previous"] for x in revenue_lines)
cogs = sum(x["current"] for x in cogs_lines)
cogs_prev = sum(x["previous"] for x in cogs_lines)
gross = revenue - cogs
gross_prev = revenue_prev - cogs_prev
sga = sum(x["current"] for x in sga_lines)
sga_prev = sum(x["previous"] for x in sga_lines)
op = gross - sga
op_prev = gross_prev - sga_prev
interest_income, interest_income_prev = 2850000, 2600000
interest_cost = sum_voucher_by_code(711)
interest_cost_prev = prev_by_label.get("이자비용", 0) or 0
pretax = op + interest_income - interest_cost
pretax_prev = op_prev + interest_income_prev - interest_cost_prev
TAX_RATE = 0.22
tax = xround(max(pretax, 0) * TAX_RATE)
tax_prev = xround(max(pretax_prev, 0) * TAX_RATE)
net = pretax - tax
net_prev = pretax_prev - tax_prev

income_statement = {
    "revenueLines": revenue_lines,
    "cogsLines": cogs_lines,
    "sgaLines": sga_lines,
    "totals": {
        "revenue": {"current": revenue, "previous": revenue_prev},
        "cogs": {"current": cogs, "previous": cogs_prev},
        "grossProfit": {"current": gross, "previous": gross_prev},
        "sga": {"current": sga, "previous": sga_prev},
        "operatingProfit": {"current": op, "previous": op_prev},
        "interestIncome": {"current": interest_income, "previous": interest_income_prev},
        "interestCost": {"current": interest_cost, "previous": interest_cost_prev},
        "pretaxProfit": {"current": pretax, "previous": pretax_prev},
        "tax": {"current": tax, "previous": tax_prev},
        "netProfit": {"current": net, "previous": net_prev},
    },
    "taxRate": TAX_RATE,
}

# ------------------------------------------------------------ 04 자금 입출금
cash_ws = wb["04_자금_입출금"]
cashflow = []
balance = 0
for r in rows(cash_ws, last_row=114):
    inflow = r["입금액"] or 0
    outflow = r["출금액"] or 0
    if r["구분"] == "기초":
        balance = r["잔액"] or 0
    else:
        balance = balance + inflow - outflow
    cashflow.append(
        {
            "date": iso(r["일자"]),
            "type": r["구분"],
            "account": r["계정"],
            "counterparty": r["거래처"],
            "memo": r["적요"],
            "inflow": inflow,
            "outflow": outflow,
            "balance": balance,
            "bank": r["은행계좌"],
        }
    )

# ------------------------------------------------------------ 05 채권 미수금
AGE_BUCKETS = ["회수완료", "정상", "30일이하", "31~60일", "61~90일", "90일초과"]
receivables = []
for r in rows(wb["05_채권_미수금"], last_row=59):
    billed = r["청구금액"] or 0
    paid = r["입금액"] or 0
    outstanding = billed - paid
    due = d(r["청구일"])
    due_date = date.fromordinal(due.toordinal() + 30)
    if outstanding <= 0:
        overdue_days = 0
        bucket = "회수완료"
    else:
        overdue_days = (BASE_DATE - due_date).days
        if overdue_days <= 0:
            bucket = "정상"
        elif overdue_days <= 30:
            bucket = "30일이하"
        elif overdue_days <= 60:
            bucket = "31~60일"
        elif overdue_days <= 90:
            bucket = "61~90일"
        else:
            bucket = "90일초과"
    status = "완료" if outstanding <= 0 else ("미회수" if paid == 0 else "부분회수")
    receivables.append(
        {
            "invoiceNo": r["청구번호"],
            "customerCode": r["거래처코드"],
            "customer": r["거래처명"],
            "industry": r["산업군"],
            "issueDate": iso(r["청구일"]),
            "paymentTerm": r["결제조건"],
            "dueDate": due_date.isoformat(),
            "billed": billed,
            "paid": paid,
            "outstanding": outstanding,
            "overdueDays": overdue_days,
            "ageBucket": bucket,
            "owner": r["담당자"],
            "status": status,
        }
    )

# --------------------------------------------------------- 06 영업 파이프라인
pipeline = []
for r in rows(wb["06_영업_파이프라인"], last_row=65):
    amount = r["예상금액"] or 0
    prob = r["확률"] or 0
    last_activity = d(r["최근활동일"])
    pipeline.append(
        {
            "dealId": r["딜ID"],
            "customer": r["거래처명"],
            "industry": r["산업군"],
            "region": r["지역"],
            "rep": r["영업담당자"],
            "team": r["부서"],
            "productGroup": r["제품군"],
            "stage": r["단계"],
            "probability": prob,
            "amount": amount,
            "weightedAmount": xround(amount * prob),
            "createdAt": iso(r["생성일"]),
            "lastActivityAt": iso(r["최근활동일"]),
            "expectedCloseAt": iso(r["예상마감일"]),
            "idleDays": (BASE_DATE - last_activity).days if last_activity else None,
            "channel": r["유입채널"],
            "competitor": r["경쟁사"],
            "note": r["비고"],
        }
    )

# --------------------------------------------------------- 07 영업 활동로그
activities = []
for r in rows(wb["07_영업_활동로그"], last_row=223):
    activities.append(
        {
            "activityId": r["활동ID"],
            "date": iso(r["일자"]),
            "rep": r["영업담당자"],
            "team": r["부서"],
            "customer": r["거래처명"],
            "dealId": r["딜ID"],
            "type": r["활동유형"],
            "durationMin": r["소요시간(분)"] or 0,
            "result": r["결과"],
            "nextAction": r["다음액션"],
            "nextActionDate": iso(r["다음액션일"]),
        }
    )

# ---------------------------------------------------------- 08 인사 직원명부
employees = []
for r in rows(wb["08_인사_직원명부"], last_row=48):
    hired = d(r["입사일"])
    employees.append(
        {
            "empNo": r["사번"],
            "name": r["성명"],
            "dept": r["부서"],
            "rank": r["직급"],
            "hiredAt": iso(r["입사일"]),
            "tenureMonths": months_between(hired, BASE_DATE),
            "employmentType": r["고용형태"],
            "annualSalary": r["연봉"] or 0,
            "email": (r["이메일"] or "").replace("@nextwave.co.kr", "@surfers.co.kr"),
            "status": r["재직상태"],
            "resignedAt": iso(r["퇴사일"]),
            "workplace": r["근무지"],
        }
    )

# ---------------------------------------------------------- 10 인사 근태현황
attendance = []
overtime_by_emp = {}
for r in rows(wb["10_인사_근태현황"], last_row=48):
    overtime = r["연장근로시간"] or 0
    scheduled = r["소정근무일수"] or 0
    annual_leave = r["연차사용일수"] or 0
    absence = r["결근일수"] or 0
    actual_days = scheduled - annual_leave - absence
    overtime_by_emp[r["사번"]] = overtime
    attendance.append(
        {
            "empNo": r["사번"],
            "name": r["성명"],
            "dept": r["부서"],
            "overtimeHours": overtime,
            "scheduledDays": scheduled,
            "actualDays": actual_days,
            "nightHours": r["야간근로시간"] or 0,
            "annualLeaveUsed": annual_leave,
            "annualLeaveLeft": r["잔여연차"] or 0,
            "lateCount": r["지각횟수"] or 0,
            "absenceDays": absence,
            "remoteDays": r["재택근무일수"] or 0,
            "totalWorkHours": actual_days * 8 + overtime,
        }
    )

# ---------------------------------------------------------- 09 인사 급여대장
payroll = []
for r in rows(wb["09_인사_급여대장"], last_row=48):
    base = r["기본급"] or 0
    meal = r["식대"] or 0
    position = r["직책수당"] or 0
    ot_hours = overtime_by_emp.get(r["사번"], 0)
    ot_pay = xround(base / 209 * 1.5 * ot_hours)
    gross_pay = base + meal + position + ot_pay
    pension = xround(min(base, 6170000) * 0.045)
    health = xround(base * 0.03545)
    care = xround(health * 0.1295)
    employment = xround(base * 0.009)
    rate = r["간이세율"] or 0
    income_tax = xround(gross_pay * rate)
    local_tax = xround(income_tax * 0.1, -1)
    deduction = pension + health + care + employment + income_tax + local_tax
    payroll.append(
        {
            "empNo": r["사번"],
            "name": r["성명"],
            "dept": r["부서"],
            "baseSalary": base,
            "mealAllowance": meal,
            "positionAllowance": position,
            "overtimePay": ot_pay,
            "grossPay": gross_pay,
            "pension": pension,
            "health": health,
            "longTermCare": care,
            "employmentIns": employment,
            "taxRate": rate,
            "incomeTax": income_tax,
            "localTax": local_tax,
            "totalDeduction": deduction,
            "netPay": gross_pay - deduction,
        }
    )

# ---------------------------------------------------------- 11 업무 프로젝트
projects = []
for r in rows(wb["11_업무_프로젝트"], last_row=23):
    budget = r["예산"] or 0
    spent = r["투입비용"] or 0
    end = d(r["종료예정일"])
    projects.append(
        {
            "code": r["프로젝트코드"],
            "name": r["프로젝트명"],
            "customer": r["고객사"],
            "pm": r["PM"],
            "dept": r["담당부서"],
            "type": r["구분"],
            "startAt": iso(r["시작일"]),
            "endAt": iso(r["종료예정일"]),
            "progress": r["진행률"] or 0,
            "status": r["상태"],
            "contractAmount": r["계약금액"] or 0,
            "budget": budget,
            "spent": spent,
            "budgetUsage": (spent / budget) if budget else 0,
            "budgetLeft": budget - spent,
            "daysLeft": (end - BASE_DATE).days if end else None,
            "risk": r["리스크등급"],
            "issue": r["이슈"],
        }
    )

# ------------------------------------------------------------ 12 업무 태스크
tasks = []
for r in rows(wb["12_업무_태스크"], last_row=97):
    due = d(r["마감일"])
    done = d(r["완료일"])
    if r["상태"] == "완료":
        delay = max(0, (done - due).days) if done and due else 0
    else:
        delay = max(0, (BASE_DATE - due).days) if due else 0
    tasks.append(
        {
            "taskId": r["태스크ID"],
            "projectCode": r["프로젝트코드"],
            "name": r["업무명"],
            "owner": r["담당자"],
            "dept": r["부서"],
            "priority": r["우선순위"],
            "status": r["상태"],
            "startAt": iso(r["시작일"]),
            "dueAt": iso(r["마감일"]),
            "doneAt": iso(r["완료일"]),
            "progress": r["진행률"] or 0,
            "estimatedHours": r["예상공수(h)"] or 0,
            "actualHours": r["실투입(h)"] or 0,
            "delayDays": delay,
        }
    )

company = {
    "name": "(주)서퍼스",
    "nameEn": "Surfers Co., Ltd.",
    "industry": "스마트물류 SaaS 플랫폼 및 산업용 IoT 하드웨어",
    "hq": "경기도 성남시 분당구 판교로 255",
    "foundedAt": "2016-03-14",
    "periodStart": "2026-07-01",
    "periodEnd": BASE_DATE.isoformat(),
    "baseDate": BASE_DATE.isoformat(),
    "currency": "KRW",
    "vatRate": VAT_RATE,
}

os.makedirs(OUT, exist_ok=True)
files = {
    "company.json": company,
    "sales.json": sales,
    "vouchers.json": vouchers,
    "income-statement.json": income_statement,
    "cashflow.json": cashflow,
    "receivables.json": receivables,
    "pipeline.json": pipeline,
    "activities.json": activities,
    "employees.json": employees,
    "payroll.json": payroll,
    "attendance.json": attendance,
    "projects.json": projects,
    "tasks.json": tasks,
}
for fname, payload in files.items():
    with open(os.path.join(OUT, fname), "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)
    print(f"{fname:26} {len(payload) if isinstance(payload, list) else '-'}")

# ------------------------------------------------------------------ 검증 출력
print("\n--- 검증 (README 기대값과 비교) ---")
print(f"매출액        {revenue:,.0f}   (기대 1,506,726,000)")
print(f"매출원가      {cogs:,.0f}   (기대 746,244,000)")
print(f"판관비        {sga:,.0f}   (기대 506,647,369)")
print(f"영업이익      {op:,.0f}   (기대 253,834,631)")
print(f"당기순이익    {net:,.0f}   (기대 193,818,012)")
print(f"매출 증감률   {(revenue-revenue_prev)/revenue_prev:.3%}  (기대 +6.7%)")
print(f"기말현금      {cashflow[-1]['balance']:,.0f}   (기대 1,965,900,000)")
tot_ar = sum(r['outstanding'] for r in receivables)
over30 = sum(r['outstanding'] for r in receivables if r['overdueDays'] > 30)
print(f"총 미수금     {tot_ar:,.0f}   (기대 1,304,630,000)")
print(f"30일초과연체  {over30:,.0f}   (기대 799,730,000)")
print(f"파이프라인    {sum(p['amount'] for p in pipeline):,.0f}  (기대 8,914,700,000)")
print(f"가중예상      {sum(p['weightedAmount'] for p in pipeline):,.0f}  (기대 3,511,095,000)")
print(f"급여지급총액  {sum(p['grossPay'] for p in payroll):,.0f}   (기대 203,547,369)")
print(f"지연 태스크   {sum(1 for t in tasks if t['delayDays']>0)}건   (기대 34건)")
print(f"미접촉 딜     {sum(1 for p in pipeline if (p['idleDays'] or 0)>14)}건   (기대 37건)")
